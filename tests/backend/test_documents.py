import asyncio
import hashlib
import json
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from threading import Event

import httpx2
import pytest
import yaml
from docx import Document
from fastapi.testclient import TestClient
from openai import AsyncOpenAI
from pydantic import SecretStr, ValidationError

from backend.app.core.errors import ApplicationError
from backend.app.documents.config import DocumentConfig, load_config
from backend.app.documents.processing import DoclingChunker
from backend.app.main import create_app

POLICY = b"# Handbook\n\n## Leave policy\n\nEmployees receive 25 days of paid leave.\n"


@pytest.fixture
def document_settings(app_settings):
    config = load_config(app_settings.docs_config_path)
    config.embeddings.dimensions = 8
    config.embeddings.max_retries = 0
    config.embeddings.batch_size = 2
    config.chunking.max_tokens = 64
    app_settings.docs_config_path.write_text(
        yaml.safe_dump(config.model_dump(mode="json")), encoding="utf-8"
    )
    return app_settings


@pytest.fixture
def embeddings_stub(monkeypatch):
    class Stub:
        requests = []
        fail = False

        async def respond(self, request):
            body = json.loads(request.content)
            self.requests.append(body)
            if self.fail:
                return httpx2.Response(500, json={"error": {"message": "private provider error"}})
            data = []
            for index, text in enumerate(body["input"]):
                digest = hashlib.sha256(text.encode()).digest()
                vector = [(digest[i % 32] + 1) / 256 for i in range(body["dimensions"])]
                data.append({"index": index, "object": "embedding", "embedding": vector})
            return httpx2.Response(
                200,
                json={
                    "object": "list",
                    "model": body["model"],
                    "data": list(reversed(data)),
                    "usage": {"prompt_tokens": 10, "total_tokens": 10},
                },
            )

    stub = Stub()

    def factory(**kwargs):
        return AsyncOpenAI(
            http_client=httpx2.AsyncClient(transport=httpx2.MockTransport(stub.respond)), **kwargs
        )

    monkeypatch.setattr("backend.app.documents.embeddings.AsyncOpenAI", factory)
    return stub


def upload(client, content=POLICY, name="policy.md"):
    return client.post("/docs/ingest", files=[("files", (name, content))])


def test_ingest_duplicates_citations_and_restart(document_settings, embeddings_stub):
    app = create_app(document_settings)
    with TestClient(app) as client:
        response = upload(client)
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["results"][0]["status"] == "ingested"
        assert body["embedded_texts"] == 1
        collection = body["collection"]
        document_id = body["results"][0]["document_id"]
        count = len(embeddings_stub.requests)
        assert upload(client).json()["results"][0]["status"] == "unchanged"
        assert upload(client, name="copy.md").json()["results"][0]["status"] == "duplicate"
        assert len(embeddings_stub.requests) == count

        response = client.post("/docs/retrieve", json={"query": "leave policy", "top_k": 3})
        assert response.status_code == 200, response.text
        hit = response.json()["hits"][0]
        assert hit["headings"] == ["Handbook", "Leave policy"]
        assert hit["pages"] == []  # Markdown has no inherent page numbers.
        assert hit["references"][0]["ref"].startswith("#/texts/")
        assert {source["filename"] for source in hit["sources"]} == {"policy.md", "copy.md"}
        active, documents = app.state.document_service.catalog.snapshot()
        assert active.collection == collection
        for doc in documents:
            assert (app.state.document_service.root / doc.original_path).read_bytes() == POLICY

    # A fresh app opens the same on-disk Chroma index and manifest.
    with TestClient(create_app(document_settings)) as restarted:
        response = restarted.post(
            "/docs/retrieve",
            json={
                "query": "leave policy",
                "document_ids": [document_id],
            },
        )
        assert response.status_code == 200
        assert response.json()["collection"] == collection
        assert len(response.json()["hits"][0]["sources"]) == 1


def test_update_hides_superseded_chunks(document_settings, embeddings_stub):
    with TestClient(create_app(document_settings)) as client:
        first = upload(client).json()
        changed = upload(client, b"# Remote work\n\nStaff can work remotely three days per week.")
        assert changed.status_code == 200
        assert changed.json()["results"][0]["status"] == "updated"
        assert changed.json()["results"][0]["document_id"] == first["results"][0]["document_id"]
        response = client.post("/docs/retrieve", json={"query": "leave", "top_k": 10})
        assert all("25 days" not in hit["text"] for hit in response.json()["hits"])
        assert "remotely" in response.json()["hits"][0]["text"]


def test_batch_supported_formats_and_failed_file(document_settings, embeddings_stub):
    word = Document()
    word.add_heading("Security", 1)
    word.add_paragraph("Keep credentials private.")
    buffer = BytesIO()
    word.save(buffer)
    with TestClient(create_app(document_settings)) as client:
        response = client.post(
            "/docs/ingest",
            files=[
                ("files", ("policy.md", POLICY)),
                ("files", ("security.docx", buffer.getvalue())),
                ("files", ("remote.html", b"<h1>Remote work</h1><p>Remote work is available.</p>")),
                ("files", ("unsupported.exe", b"not a document")),
            ],
        )
        assert response.status_code == 200, response.text
        assert [item["status"] for item in response.json()["results"]] == [
            "ingested",
            "ingested",
            "ingested",
            "failed",
        ]
        assert response.json()["embedded_texts"] >= 3
        assert all(len(request["input"]) <= 2 for request in embeddings_stub.requests)


def test_reindex_reuses_embeddings_and_changes_settings(document_settings, embeddings_stub):
    app = create_app(document_settings)
    with TestClient(app) as client:
        first = upload(client).json()
        calls = len(embeddings_stub.requests)
        response = client.post("/docs/reindex", json={})
        assert response.status_code == 200, response.text
        assert response.json()["collection"] != first["collection"]
        assert response.json()["previous_collection"] == first["collection"]
        assert response.json()["embedded_texts"] == 0
        assert response.json()["cached_texts"] == 1
        assert len(embeddings_stub.requests) == calls

        config = load_config(document_settings.docs_config_path)
        config.chroma.hnsw.ef_search = 150
        document_settings.docs_config_path.write_text(
            yaml.safe_dump(config.model_dump(mode="json")), encoding="utf-8"
        )
        assert client.post("/docs/reindex").status_code == 200
        assert len(embeddings_stub.requests) == calls
        config.embeddings.dimensions = 12
        document_settings.docs_config_path.write_text(
            yaml.safe_dump(config.model_dump(mode="json")), encoding="utf-8"
        )
        assert client.post("/docs/reindex").json()["embedded_texts"] == 1
        assert embeddings_stub.requests[-1]["dimensions"] == 12
        assert client.post("/docs/retrieve", json={"query": "leave"}).status_code == 200
        assert embeddings_stub.requests[-1]["dimensions"] == 12


def test_failed_reindex_leaves_active_index_searchable(
    document_settings, embeddings_stub, monkeypatch
):
    app = create_app(document_settings)
    with TestClient(app) as client:
        collection = upload(client).json()["collection"]
        vectors = app.state.document_service.vectors
        original = vectors.upsert

        def fail_write(*args):
            raise RuntimeError("simulated failed staging write")

        monkeypatch.setattr(vectors, "upsert", fail_write)
        assert client.post("/docs/reindex").status_code == 500
        assert app.state.document_service.catalog.snapshot()[0].collection == collection
        assert len(vectors.client.list_collections()) == 1
        monkeypatch.setattr(vectors, "upsert", original)
        response = client.post("/docs/retrieve", json={"query": "leave"})
        assert response.status_code == 200
        assert response.json()["collection"] == collection
        assert response.json()["hits"]


def test_queries_continue_during_reindex(document_settings, embeddings_stub, monkeypatch):
    app = create_app(document_settings)
    started, release = Event(), Event()
    with TestClient(app) as client:
        old = upload(client).json()["collection"]
        vectors = app.state.document_service.vectors
        original = vectors.upsert

        def slow_write(*args):
            started.set()
            assert release.wait(10)
            return original(*args)

        monkeypatch.setattr(vectors, "upsert", slow_write)
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(client.post, "/docs/reindex")
            try:
                assert started.wait(10)
                response = client.post("/docs/retrieve", json={"query": "leave"})
                assert response.status_code == 200
                assert response.json()["collection"] == old
                assert response.json()["hits"]
            finally:
                release.set()
            assert future.result(timeout=10).status_code == 200


def test_failed_embeddings_do_not_publish_updates(document_settings, embeddings_stub):
    app = create_app(document_settings)
    with TestClient(app) as client:
        upload(client)
        before = app.state.document_service.catalog.snapshot()
        embeddings_stub.fail = True
        response = upload(client, b"# New document\n\nThis update must not replace the old index.")
        assert response.status_code == 502
        assert app.state.document_service.catalog.snapshot() == before


def test_ingest_limits_validation_and_empty_retrieval(document_settings, embeddings_stub):
    config = load_config(document_settings.docs_config_path)
    config.ingestion.max_file_mb = 0.001
    document_settings.docs_config_path.write_text(yaml.safe_dump(config.model_dump(mode="json")))
    with TestClient(create_app(document_settings)) as client:
        assert client.post("/docs/retrieve", json={"query": "unknown"}).json()["hits"] == []
        assert embeddings_stub.requests == []
        assert upload(client, b"x" * 2000).status_code == 413
        assert client.post("/docs/retrieve", json={"query": "   "}).status_code == 422
        assert client.post("/docs/retrieve", json={"query": "x", "top_k": 9999}).status_code == 422
        duplicate_sources = client.post(
            "/docs/ingest",
            files=[
                ("files", ("same.md", POLICY)),
                ("files", ("same.md", POLICY)),
            ],
        )
        assert duplicate_sources.status_code == 422


def test_chunker_retains_page_numbers_and_bounds():
    from docling_core.types.doc import (
        BoundingBox,
        DocItemLabel,
        DoclingDocument,
        ProvenanceItem,
        Size,
    )

    document = DoclingDocument(name="paged-policy")
    document.add_page(page_no=3, size=Size(width=600, height=800))
    document.add_heading("Leave policy")
    text = "Employees have paid leave."
    document.add_text(
        label=DocItemLabel.TEXT,
        text=text,
        prov=ProvenanceItem(
            page_no=3, bbox=BoundingBox(l=10, t=20, r=400, b=60), charspan=(0, len(text))
        ),
    )
    chunks = DoclingChunker().chunk(document, DocumentConfig())
    assert chunks[0].pages == [3]
    assert chunks[0].headings == ["Leave policy"]
    provenance = chunks[0].references[0]["provenance"][0]
    assert provenance["bbox"]["l"] == 10
    assert provenance["charspan"] == [0, len(text)]


def test_embedding_batches_and_cross_document_cache(document_settings, embeddings_stub):
    app = create_app(document_settings)
    with TestClient(app):
        service = app.state.document_service
        config = service.config.embeddings.model_copy(
            update={
                "batch_size": 2,
                "max_input_tokens": 100,
                "max_batch_tokens": 110,
            }
        )
        texts = ["policy " * 50, "security " * 50, "remote " * 50]
        first = asyncio.run(service.embedder.embed(texts + texts[:1], config, "cl100k_base"))
        assert first.embedded == 3
        assert len(embeddings_stub.requests) == 2
        assert first.vectors[0] == first.vectors[3]
        second = asyncio.run(service.embedder.embed(texts, config, "cl100k_base"))
        assert second.embedded == 0 and second.cached == 3
        assert second.vectors == first.vectors[:3]

        # Cached vectors must still respect newly tightened input limits.
        smaller = config.model_copy(update={"max_input_tokens": 20})
        with pytest.raises(ApplicationError, match="token limit"):
            asyncio.run(service.embedder.embed(texts, smaller, "cl100k_base"))


def test_partial_vector_write_cannot_publish_update(
    document_settings, embeddings_stub, monkeypatch
):
    app = create_app(document_settings)
    with TestClient(app) as client:
        upload(client)
        vectors = app.state.document_service.vectors
        original = vectors.upsert

        def write_then_fail(*args):
            original(*args)
            raise RuntimeError("failed before catalog commit")

        monkeypatch.setattr(vectors, "upsert", write_then_fail)
        changed = upload(client, b"# Changed policy\n\nNo paid leave is allowed.")
        assert changed.status_code == 500
        response = client.post("/docs/retrieve", json={"query": "leave", "top_k": 20})
        assert response.status_code == 200
        assert all("25 days" in hit["text"] for hit in response.json()["hits"])


def test_missing_key_preserves_original_without_publishing(document_settings):
    settings = document_settings.model_copy(update={"openai_api_key": SecretStr("")})
    app = create_app(settings)
    with TestClient(app) as client:
        response = upload(client)
        assert response.status_code == 503
        assert app.state.document_service.catalog.snapshot() == (None, [])
        assert list((app.state.document_service.root / "originals").glob("*.md"))


def test_forced_reindex_and_bad_config(document_settings, embeddings_stub):
    app = create_app(document_settings)
    with TestClient(app) as client:
        upload(client)
        result = client.post("/docs/reindex", json={"force_reparse": True, "force_reembed": True})
        assert result.status_code == 200
        assert result.json()["embedded_texts"] == 1
        active = app.state.document_service.catalog.snapshot()[0]
        document_settings.docs_config_path.write_text("parsing: [invalid yaml", encoding="utf-8")
        assert client.post("/docs/reindex").status_code == 422
        assert app.state.document_service.catalog.snapshot()[0] == active


def test_config_parameters_are_validated_and_documented():
    from shared.settings import PROJECT_ROOT

    config = load_config(PROJECT_ROOT / "config.yaml")
    # config.yaml is intentionally user-tunable; it need not equal every library default.
    documentation = (PROJECT_ROOT / "backend/app/documents/README.md").read_text(encoding="utf-8")
    documentation += (PROJECT_ROOT / "backend/app/evaluation/README.md").read_text(encoding="utf-8")
    documentation += (PROJECT_ROOT / "backend/app/investment/README.md").read_text(encoding="utf-8")

    def check_fields(value, prefix=""):
        for key, child in value.items():
            name = f"{prefix}.{key}" if prefix else key
            map_fields = {
                "parsing.format_map",
                "investment.profiles",
                "investment.dcf.scenarios",
                "investment.fundamentals_bands",
                "investment.trend_bands",
                "investment.business_risk_bands",
                "investment.market_risk_bands",
                "investment.valuation_bands",
            }
            if isinstance(child, dict) and name not in map_fields:
                check_fields(child, name)
            else:
                assert f"`{name}`" in documentation

    check_fields(config.model_dump(mode="json"))
    for invalid in (
        {"chunking": {"tokenizer": "missing-encoding"}},
        {"parsing": {"format_map": {"pdf": "missing-format"}}},
        {"embeddings": {"dimensions": 2048}},
        {"retrieval": {"top_k": 100, "max_top_k": 10}},
        {"unknown_setting": True},
    ):
        with pytest.raises(ValidationError):
            DocumentConfig.model_validate(invalid)


@pytest.mark.parametrize("category", ["timeout", "unknown"])
def test_docling_failure_logs_category_and_reports_timeouts(tmp_path, monkeypatch, category):
    import io
    import logging
    from types import SimpleNamespace

    from docling.datamodel.base_models import ConversionStatus, ErrorItem
    from docling.document_converter import DocumentConverter

    from backend.app.documents.processing import DoclingParser
    from shared.logging_config import JsonFormatter

    result = SimpleNamespace(
        status=ConversionStatus.PARTIAL_SUCCESS,
        errors=[
            ErrorItem(
                component_type="pipeline",
                module_name="StandardPdfPipeline",
                error_message="private-document-content",
                category=category,
                page_no=2,
            )
        ],
    )
    monkeypatch.setattr(DocumentConverter, "convert", lambda *args, **kwargs: result)
    original = tmp_path / "original.pdf"
    original.write_bytes(b"dummy-pdf")
    output = io.StringIO()
    handler = logging.StreamHandler(output)
    handler.setFormatter(JsonFormatter("backend"))
    logger = logging.getLogger("credicorp.backend")
    logger.addHandler(handler)
    try:
        with pytest.raises(ApplicationError) as failure:
            DoclingParser().parse(original, "pdf", DocumentConfig())
    finally:
        logger.removeHandler(handler)
        handler.close()
    assert "private-document-content" not in str(failure.value) + output.getvalue()
    if category == "timeout":
        assert "timed out" in failure.value.detail
        assert "180s" in failure.value.detail
    else:
        assert "partial_success" in failure.value.detail
    records = [json.loads(line) for line in output.getvalue().splitlines()]
    assert records[0]["message"] == "docling_conversion_started"
    assert records[-1]["message"] == "docling_conversion_failed"
    assert records[-1]["conversion_errors"][0]["category"] == category
    assert records[-1]["conversion_errors"][0]["page_no"] == 2
